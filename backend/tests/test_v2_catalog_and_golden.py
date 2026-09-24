"""V2 catalogue and Golden Case regression tests.

Golden documents are reference material only; generators receive workbook
patient answers and the explicit TEST_ONLY candidate profile.
"""
from pathlib import Path
import pytest
try:
    from openpyxl import load_workbook
except ImportError:  # optional developer dependency; API runtime does not need it
    load_workbook = None
    pytestmark = pytest.mark.skip(reason="openpyxl未安装，跳过基于合成病例工作簿的回归")

from app.v2_engine import build_v2_plan, candidate_test_profile
from app.rules import assess_payload
from app.v2_knowledge import structured_catalog
from app.plan_validator import validate_plan

ROOT = Path(__file__).resolve().parents[2]
BOOK = ROOT / "docs" / "v2_upgrade_20260912" / "template" / "Q1-Q56_合成患者测试数据_A-F分型_18例.xlsx"

def _payload(case_id: str):
    ws = load_workbook(BOOK, data_only=True)["Q1-Q56宽表"]
    row = next(r for r in ws.iter_rows(min_row=3, values_only=True) if r[0] == case_id)
    q = lambda n: row[3 + n]
    return {
        "sex": q(2), "height_cm": q(14), "weight_kg": q(15), "q14_height": q(14), "q15_weight": q(15),
        "q18_bodyFat": q(18), "q24_skeletalMuscle": q(24), "q26_smi": q(26), "q27_walkTest": q(27), "q6_surgeryWindow": q(6),
        "q12_respiratorySymptoms": [q(12)] if q(12) else [], "q28_weightChange": q(28), "q29_appetite": q(29),
        "q30_intake": q(30), "q31_eatingDifficulties": [q(31)] if q(31) else [], "q32_foodAllergy": [q(32)] if q(32) else [],
        "q34_dietPattern": [q(34)] if q(34) else [], "q43_metabolicConditions": [q(43)] if q(43) else [],
        "q51_executionBarriers": [q(51)] if q(51) else [], "q56_goal": q(56), "age_years": 54,
    }

def test_v2_catalog_has_all_structured_identifiers():
    catalog = structured_catalog()
    assert len(catalog["FOOD"]) == 49
    assert len(catalog["EXERCISE"]) == 31
    assert {x["pulmonary_id"] for x in catalog["PULMONARY"]} == {f"P0{i}" for i in range(1, 7)}
    for item in catalog["FOOD"]:
        assert item["component_id"] and item["ingredients"] and item["ingredient_amounts"]
        assert item["raw_or_cooked_basis"]
        # V1.7 is the confirmed runtime FOOD catalogue; legacy V2 catalogue
        # entries remain DRAFT.  Both statuses are valid for this structural
        # test, while governance promotion is handled separately.
        assert item["status"] in {"ACTIVE", "DRAFT"}
    for item in catalog["EXERCISE"]:
        assert item["exercise_id"] and item["name"] and item["steps"]
    for item in catalog["PULMONARY"]:
        assert item["pulmonary_id"] and item["name"] and item["steps"]

def test_candidate_generates_from_answers_not_golden_files():
    payload = _payload("SYN-A01")
    evaluation = assess_payload(payload)
    plan = build_v2_plan(payload, evaluation, active_configs=candidate_test_profile())
    assert evaluation["phenotype_code"] == "A"
    assert plan["goal_source"] == "SYSTEM_DEFAULT_AF"
    assert plan["energy_calculation"]["ree_source"] == "P2_APPROVED_FORMULA"
    assert len(plan["diet_plan"]["components"]) == 4
    assert plan["exercise_plan"] and plan["pulmonary_prehab_plan"]
    assert all(x.get("component_id") is None or x.get("component_id") for x in plan["diet_plan"]["components"])

def test_all_18_workbook_phenotypes_match_expected():
    ws = load_workbook(BOOK, data_only=True)["Q1-Q56宽表"]
    results = []
    for row in ws.iter_rows(min_row=3, values_only=True):
        if not row[0] or not str(row[0]).startswith("SYN-"):
            continue
        actual = assess_payload(_payload(row[0]))["phenotype_code"]
        results.append((row[0], row[1], actual))
    assert len(results) == 18
    assert all(expected == actual for _, expected, actual in results)

def test_f03_safety_overrides_candidate_profile():
    payload = _payload("SYN-F03")
    evaluation = assess_payload(payload)
    # Workbook Q56 is encoded in Chinese text; assessment parser preserves it.
    plan = build_v2_plan(payload, evaluation, active_configs=candidate_test_profile())
    assert evaluation["phenotype_code"] == "F"
    assert evaluation["safety"] == "red"
    assert plan["enhanced_eligible"] is False
    assert plan["goal_conflict"] is True
    assert plan["publication_blocked"] is True
    assert plan["manual_review_required"] is True
    assert plan["exercise_plan"] == [] and plan["pulmonary_prehab_plan"] == []

def test_plan_validator_blocks_pending_or_red_plan():
    payload = _payload("SYN-A01")
    evaluation = assess_payload(payload)
    plan = build_v2_plan(payload, evaluation)
    result = validate_plan(plan)
    assert result["status"] == "FAIL"
